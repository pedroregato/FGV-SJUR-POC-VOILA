import pandas as pd
import re
import os
from datetime import datetime
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


def aplicar_mascara_cnj(numero: str) -> str:
    """Aplica a máscara CNJ: NNNNNNN-DD.AAAA.J.TR.OOOO"""
    numero = re.sub(r"[^\d]", "", str(numero))
    if len(numero) != 20:
        return numero
    return f"{numero[0:7]}-{numero[7:9]}.{numero[9:13]}.{numero[13]}.{numero[14:16]}.{numero[16:20]}"


def ajustar_planilha(resultados_path, gerpro_path, output_path=None):
    # Carrega planilha de resultados
    df_res = pd.read_excel(resultados_path)

    # Cria coluna mascarada
    df_res.insert(
        loc=1,
        column="numero_processo_mascarado",
        value=df_res["numero_processo"].apply(aplicar_mascara_cnj)
    )

    # Carrega GERPRO (colunas A e B)
    df_gerpro = pd.read_excel(gerpro_path, usecols=[0, 1], header=0)
    df_gerpro.columns = ["coluna_A_gerpro", "numero_na_justica"]
    df_gerpro["numero_processo_mascarado"] = df_gerpro["numero_na_justica"].apply(aplicar_mascara_cnj)

    total_registros_gerpro = len(df_gerpro)

    # Merge
    df_final = pd.merge(
        df_res,
        df_gerpro[["coluna_A_gerpro", "numero_processo_mascarado"]],
        on="numero_processo_mascarado",
        how="left"
    )
    df_final.rename(columns={"coluna_A_gerpro": "coluna_A_Gerpro"}, inplace=True)

    # Caminho de saída
    if not output_path:
        base, fname = os.path.split(resultados_path)
        output_path = os.path.join(base, f"ajustado_{fname}")

    # --- RESUMO EXECUTIVO ---
    resumo = {}

    resumo["Total de registros na planilha origem do GERPRO"] = total_registros_gerpro
    resumo["Total de registros encontrados pela API"] = (df_final["encontrado_na_api"] == "Sim").sum()
    resumo["Total de registros não encontrados pela API"] = (df_final["encontrado_na_api"] == "Não").sum()
    resumo["Total de registros sem correspondência no GERPRO"] = df_final["coluna_A_Gerpro"].isna().sum()
    resumo["Total de registros com Trânsito em Julgado"] = (df_final["tem_transito_em_julgado"] == "Sim").sum()

    # --- Ajuste de datas ---
    hoje = datetime.now()

    ajuizamentos = pd.to_datetime(df_final["data_ajuizamento"], errors="coerce")
    transitos = pd.to_datetime(df_final["data_transito_em_julgado"], errors="coerce")

    ajuiz_validos = ajuizamentos[(ajuizamentos.dt.year >= 1950) & (ajuizamentos <= hoje)]
    transitos_validos = transitos[(transitos.notna()) & (transitos <= hoje)]

    descartados_ajuiz = ajuizamentos.notna().sum() - ajuiz_validos.notna().sum()
    descartados_transito = transitos.notna().sum() - transitos_validos.notna().sum()

    resumo["Menor data de ajuizamento"] = ajuiz_validos.min() if not ajuiz_validos.empty else "N/A"
    resumo["Maior data de ajuizamento"] = ajuiz_validos.max() if not ajuiz_validos.empty else "N/A"
    resumo["Menor data do trânsito em julgado"] = transitos_validos.min() if not transitos_validos.empty else "N/A"
    resumo["Maior data do trânsito em julgado"] = transitos_validos.max() if not transitos_validos.empty else "N/A"

    resumo["Registros descartados (datas inválidas ajuizamento)"] = descartados_ajuiz
    resumo["Registros descartados (datas inválidas trânsito)"] = descartados_transito

    df_resumo = pd.DataFrame(list(resumo.items()), columns=["Indicador", "Valor"])

    # --- SALVA EXCEL ---
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_final.to_excel(writer, index=False, sheet_name="Ajustado")
        df_resumo.to_excel(writer, index=False, sheet_name="Resumo")

        workbook = writer.book
        ws_resumo = workbook["Resumo"]

        # Ajusta largura automática das colunas do resumo
        for col in ws_resumo.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws_resumo.column_dimensions[col_letter].width = min(max_length + 2, 50)

        # --- Gráfico por Tribunal ---
        if "tribunal" in df_final.columns:
            df_tribunal = df_final.groupby("tribunal").agg(
                total_registros=("numero_processo", "count"),
                transitos=("tem_transito_em_julgado", lambda x: (x == "Sim").sum())
            ).reset_index()

            start_row = ws_resumo.max_row + 3
            for r_idx, row in enumerate([df_tribunal.columns.tolist()] + df_tribunal.values.tolist(), start_row):
                for c_idx, value in enumerate(row, 1):
                    ws_resumo.cell(row=r_idx, column=c_idx, value=value)

            chart = BarChart()
            chart.title = "Totais e Trânsitos em Julgado por Tribunal"
            chart.y_axis.title = "Quantidade"
            chart.x_axis.title = "Tribunal"

            data = Reference(ws_resumo, min_col=2, max_col=3, min_row=start_row, max_row=start_row + len(df_tribunal))
            cats = Reference(ws_resumo, min_col=1, min_row=start_row + 1, max_row=start_row + len(df_tribunal))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)

            ws_resumo.add_chart(chart, f"E{start_row}")

        # --- Gráfico por Classe ---
        if "classe" in df_final.columns:
            df_classe = df_final.groupby("classe").agg(
                total_registros=("numero_processo", "count"),
                transitos=("tem_transito_em_julgado", lambda x: (x == "Sim").sum())
            ).reset_index()

            # ✅ remove classes com zero em ambos
            df_classe = df_classe[(df_classe["total_registros"] > 0) | (df_classe["transitos"] > 0)]

            start_row = ws_resumo.max_row + 3
            for r_idx, row in enumerate([df_classe.columns.tolist()] + df_classe.values.tolist(), start_row):
                for c_idx, value in enumerate(row, 1):
                    ws_resumo.cell(row=r_idx, column=c_idx, value=value)

            chart = BarChart()
            chart.title = "Totais e Trânsitos em Julgado por Classe"
            chart.y_axis.title = "Quantidade"
            chart.x_axis.title = "Classe"

            data = Reference(ws_resumo, min_col=2, max_col=3, min_row=start_row, max_row=start_row + len(df_classe))
            cats = Reference(ws_resumo, min_col=1, min_row=start_row + 1, max_row=start_row + len(df_classe))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)

            ws_resumo.add_chart(chart, f"E{start_row}")

    print(f"✅ Planilha ajustada e resumo salvo em: {output_path}")
    return df_final, df_resumo


if __name__ == "__main__":
    resultados_path = r"F:\FGV-SJUR\resultado_transito_julgado\processos_transito_julgado_20250930_160155.xlsx"
    gerpro_path = r"F:\FGV-SJUR\sjur-poc-voila\data_analysis\Gerpro_Processos_Ativos_2025-09-26-20h39m.xlsx"
    ajustar_planilha(resultados_path, gerpro_path)
